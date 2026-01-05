"""
Knowledge Base System for Roca3D
Provides persistent storage and research paper ingestion capabilities.
"""

import os
import json
import sqlite3
import hashlib
import time
from datetime import datetime
from typing import Dict, List, Any, Optional, Tuple, Set
from dataclasses import dataclass, field
from enum import Enum
import threading
import queue
import re

# Document processing libraries
try:
    import PyPDF2
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import nltk
    from nltk.tokenize import sent_tokenize, word_tokenize
    from nltk.corpus import stopwords
    from nltk.stem import WordNetLemmatizer
    NLTK_AVAILABLE = True
except ImportError:
    NLTK_AVAILABLE = False

# PyQt6 imports
from PyQt6.QtWidgets import QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton, QTextEdit, QLineEdit, QListWidget, QListWidgetItem, QProgressBar, QComboBox, QGroupBox, QSplitter, QSizePolicy
from PyQt6.QtCore import QObject, pyqtSignal, QTimer, QThread, QUrl
from PyQt6.QtGui import QDesktopServices

class DocumentType(Enum):
    """Types of documents that can be ingested"""
    PDF = "pdf"
    DOCX = "docx"
    TXT = "txt"
    HTML = "html"
    MARKDOWN = "md"
    RESEARCH_PAPER = "research_paper"
    BOOK_CHAPTER = "book_chapter"
    ARTICLE = "article"

class KnowledgeType(Enum):
    """Types of knowledge stored"""
    CONCEPT = "concept"
    THEOREM = "theorem"
    ALGORITHM = "algorithm"
    DEFINITION = "definition"
    EXAMPLE = "example"
    PROOF = "proof"
    APPLICATION = "application"
    REFERENCE = "reference"

@dataclass
class KnowledgeEntry:
    """A single entry in the knowledge base"""
    id: str
    title: str
    content: str
    knowledge_type: KnowledgeType
    document_type: DocumentType
    source_file: str
    page_number: Optional[int] = None
    section: Optional[str] = None
    tags: List[str] = field(default_factory=list)
    metadata: Dict[str, Any] = field(default_factory=dict)
    created_at: float = field(default_factory=time.time)
    updated_at: float = field(default_factory=time.time)
    embedding: Optional[List[float]] = None  # For semantic search

@dataclass
class DocumentMetadata:
    """Metadata for ingested documents"""
    filename: str
    filepath: str
    document_type: DocumentType
    title: Optional[str] = None
    authors: List[str] = field(default_factory=list)
    abstract: Optional[str] = None
    keywords: List[str] = field(default_factory=list)
    publication_date: Optional[str] = None
    doi: Optional[str] = None
    pages: int = 0
    ingested_at: float = field(default_factory=time.time)

class DocumentProcessor(QThread):
    """Background thread for processing documents"""

    processing_complete = pyqtSignal(DocumentMetadata, list)  # metadata, entries
    processing_progress = pyqtSignal(str, float)  # status, progress
    processing_error = pyqtSignal(str)

    def __init__(self, filepath: str):
        super().__init__()
        self.filepath = filepath
        self.is_cancelled = False

    def run(self):
        """Process the document and extract knowledge entries"""
        try:
            self.processing_progress.emit("Starting document processing...", 0.0)

            # Determine document type
            filename = os.path.basename(self.filepath)
            ext = os.path.splitext(filename)[1].lower()

            if ext == '.pdf':
                doc_type = DocumentType.PDF
                text_content = self._extract_pdf_text()
            elif ext == '.docx':
                doc_type = DocumentType.DOCX
                text_content = self._extract_docx_text()
            elif ext == '.txt':
                doc_type = DocumentType.TXT
                text_content = self._extract_txt_text()
            else:
                raise ValueError(f"Unsupported file type: {ext}")

            self.processing_progress.emit("Extracting metadata...", 0.2)

            # Extract metadata
            metadata = self._extract_metadata(filename, doc_type, text_content)

            self.processing_progress.emit("Processing content...", 0.5)

            # Process content into knowledge entries
            entries = self._process_content(text_content, metadata)

            self.processing_progress.emit("Finalizing...", 0.9)

            if not self.is_cancelled:
                self.processing_complete.emit(metadata, entries)

        except Exception as e:
            self.processing_error.emit(f"Processing failed: {str(e)}")

    def cancel(self):
        """Cancel processing"""
        self.is_cancelled = True

    def _extract_pdf_text(self) -> str:
        """Extract text from PDF file"""
        if not PDF_AVAILABLE:
            raise ImportError("PyPDF2 not available for PDF processing")

        text = ""
        with open(self.filepath, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"

        return text

    def _extract_docx_text(self) -> str:
        """Extract text from DOCX file"""
        if not DOCX_AVAILABLE:
            raise ImportError("python-docx not available for DOCX processing")

        doc = docx.Document(self.filepath)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"

        return text

    def _extract_txt_text(self) -> str:
        """Extract text from TXT file"""
        with open(self.filepath, 'r', encoding='utf-8', errors='ignore') as file:
            return file.read()

    def _extract_metadata(self, filename: str, doc_type: DocumentType, content: str) -> DocumentMetadata:
        """Extract metadata from document content"""
        metadata = DocumentMetadata(
            filename=filename,
            filepath=self.filepath,
            document_type=doc_type
        )

        # Try to extract title (usually first non-empty line)
        lines = content.split('\n')
        for line in lines[:10]:  # Check first 10 lines
            line = line.strip()
            if line and len(line) > 10 and not line.isupper():  # Avoid all-caps headers
                metadata.title = line
                break

        # Try to extract abstract
        content_lower = content.lower()
        abstract_match = re.search(r'abstract[:\s]*(.*?)(?:\n\n|\n[A-Z]|$)', content_lower, re.DOTALL)
        if abstract_match:
            metadata.abstract = abstract_match.group(1).strip()

        # Extract keywords
        keyword_patterns = [
            r'keywords?[:\s]*(.*?)(?:\n|\r|$)',
            r'key words?[:\s]*(.*?)(?:\n|\r|$)',
            r'subject[:\s]*(.*?)(?:\n|\r|$)'
        ]

        for pattern in keyword_patterns:
            match = re.search(pattern, content_lower, re.IGNORECASE)
            if match:
                keywords_text = match.group(1).strip()
                # Split by common delimiters
                keywords = re.split(r'[,;]\s*', keywords_text)
                metadata.keywords.extend([k.strip() for k in keywords if k.strip()])
                break

        return metadata

    def _process_content(self, content: str, metadata: DocumentMetadata) -> List[KnowledgeEntry]:
        """Process content into knowledge entries"""
        entries = []

        # Split into sections/paragraphs
        sections = self._split_into_sections(content)

        for i, section in enumerate(sections):
            if self.is_cancelled:
                break

            # Analyze section content
            section_type = self._classify_section(section['content'])

            if section_type:
                entry = KnowledgeEntry(
                    id=hashlib.md5(f"{metadata.filename}_{i}".encode()).hexdigest(),
                    title=section['title'] or f"Section {i+1}",
                    content=section['content'],
                    knowledge_type=section_type,
                    document_type=metadata.document_type,
                    source_file=metadata.filename,
                    section=section['title'],
                    tags=self._extract_tags(section['content']),
                    metadata={'section_number': i + 1}
                )
                entries.append(entry)

        return entries

    def _split_into_sections(self, content: str) -> List[Dict[str, str]]:
        """Split content into logical sections"""
        sections = []
        lines = content.split('\n')

        current_section = {'title': None, 'content': ''}
        current_title = None

        for line in lines:
            line = line.strip()
            if not line:
                continue

            # Check if line looks like a section header
            if self._is_section_header(line):
                # Save previous section if it has content
                if current_section['content'].strip():
                    sections.append(current_section)

                # Start new section
                current_section = {
                    'title': line,
                    'content': ''
                }
            else:
                current_section['content'] += line + '\n'

        # Add final section
        if current_section['content'].strip():
            sections.append(current_section)

        return sections

    def _is_section_header(self, line: str) -> bool:
        """Check if a line looks like a section header"""
        # Headers are typically short, title case, and may end with punctuation
        if len(line) > 100:  # Too long to be a header
            return False

        # Check for common header patterns
        if re.match(r'^\d+(\.\d+)*\s+[A-Z]', line):  # Numbered sections
            return True

        if line.istitle() and len(line.split()) <= 10:  # Title case, reasonable length
            return True

        # Common section words
        section_words = ['introduction', 'abstract', 'conclusion', 'methodology',
                        'results', 'discussion', 'references', 'acknowledgments']
        first_word = line.lower().split()[0] if line.split() else ""
        if first_word in section_words:
            return True

        return False

    def _classify_section(self, content: str) -> Optional[KnowledgeType]:
        """Classify the type of knowledge in a section"""
        content_lower = content.lower()

        # Check for theorems
        if re.search(r'\btheorem\b|\blemma\b|\bcorollary\b|\bproposition\b', content_lower):
            return KnowledgeType.THEOREM

        # Check for definitions
        if re.search(r'\bdefinition\b|\bdefined as\b|\bdenotes\b', content_lower):
            return KnowledgeType.DEFINITION

        # Check for algorithms
        if re.search(r'\balgorithm\b|\bpseudocode\b|\bstep \d+\b', content_lower):
            return KnowledgeType.ALGORITHM

        # Check for proofs
        if re.search(r'\bproof\b|\btherefore\b|\bthus\b|\bhence\b', content_lower):
            return KnowledgeType.PROOF

        # Check for examples
        if re.search(r'\bexample\b|\bfor instance\b|\bconsider\b', content_lower):
            return KnowledgeType.EXAMPLE

        # Default to concept
        if len(content.strip()) > 50:  # Substantial content
            return KnowledgeType.CONCEPT

        return None

    def _extract_tags(self, content: str) -> List[str]:
        """Extract relevant tags from content"""
        tags = []

        # Common technical terms that might be tags
        technical_terms = [
            'capsule network', 'convolutional neural network', 'deep learning',
            'machine learning', 'computer vision', 'artificial intelligence',
            'neural network', 'routing', 'squashing', 'pose estimation',
            'dynamic routing', 'attention mechanism', 'transformer'
        ]

        content_lower = content.lower()
        for term in technical_terms:
            if term in content_lower:
                tags.append(term)

        return tags

class KnowledgeBase:
    """Main knowledge base system with persistent storage"""

    def __init__(self, db_path: str = "knowledge_base.db"):
        self.db_path = db_path
        self._init_database()
        self.processors = {}  # Active document processors

    def _init_database(self):
        """Initialize the SQLite database"""
        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()

            # Create tables
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS documents (
                    filename TEXT PRIMARY KEY,
                    filepath TEXT,
                    document_type TEXT,
                    title TEXT,
                    authors TEXT,
                    abstract TEXT,
                    keywords TEXT,
                    publication_date TEXT,
                    doi TEXT,
                    pages INTEGER,
                    ingested_at REAL
                )
            ''')

            cursor.execute('''
                CREATE TABLE IF NOT EXISTS knowledge_entries (
                    id TEXT PRIMARY KEY,
                    title TEXT,
                    content TEXT,
                    knowledge_type TEXT,
                    document_type TEXT,
                    source_file TEXT,
                    page_number INTEGER,
                    section TEXT,
                    tags TEXT,
                    metadata TEXT,
                    created_at REAL,
                    updated_at REAL,
                    embedding TEXT
                )
            ''')

            # Create indexes for better query performance
            cursor.execute('CREATE INDEX IF NOT EXISTS idx_knowledge_type ON knowledge_entries(knowledge_type)')
            cursor.execute('CREATE INDEX IF NOT EXISTS idx_tags ON knowledge_entries(tags)')
            cursor.execute('CREATE INDEX IF NOT EXISTS idx_source_file ON knowledge_entries(source_file)')

            conn.commit()

    def ingest_document(self, filepath: str) -> str:
        """Ingest a document and return processing job ID"""
        processor = DocumentProcessor(filepath)
        job_id = hashlib.md5(f"{filepath}_{time.time()}".encode()).hexdigest()

        processor.processing_complete.connect(
            lambda metadata, entries: self._on_processing_complete(job_id, metadata, entries)
        )
        processor.processing_error.connect(
            lambda error: self._on_processing_error(job_id, error)
        )

        self.processors[job_id] = processor
        processor.start()

        return job_id

    def search_knowledge(self, query: str, knowledge_types: List[KnowledgeType] = None,
                        tags: List[str] = None, limit: int = 50) -> List[KnowledgeEntry]:
        """Search the knowledge base"""
        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()

            # Build query
            sql = "SELECT * FROM knowledge_entries WHERE 1=1"
            params = []

            if knowledge_types:
                placeholders = ','.join('?' * len(knowledge_types))
                sql += f" AND knowledge_type IN ({placeholders})"
                params.extend([kt.value for kt in knowledge_types])

            if tags:
                tag_conditions = []
                for tag in tags:
                    tag_conditions.append("tags LIKE ?")
                    params.append(f'%{tag}%')
                sql += " AND (" + " OR ".join(tag_conditions) + ")"

            # Simple text search in title and content
            if query:
                sql += " AND (title LIKE ? OR content LIKE ?)"
                params.extend([f'%{query}%', f'%{query}%'])

            sql += " ORDER BY created_at DESC LIMIT ?"
            params.append(limit)

            cursor.execute(sql, params)
            rows = cursor.fetchall()

            entries = []
            for row in rows:
                # Parse JSON fields
                tags_list = json.loads(row[8]) if row[8] else []
                metadata_dict = json.loads(row[9]) if row[9] else {}
                embedding_list = json.loads(row[12]) if row[12] else None

                entry = KnowledgeEntry(
                    id=row[0],
                    title=row[1],
                    content=row[2],
                    knowledge_type=KnowledgeType(row[3]),
                    document_type=DocumentType(row[4]),
                    source_file=row[5],
                    page_number=row[6],
                    section=row[7],
                    tags=tags_list,
                    metadata=metadata_dict,
                    created_at=row[10],
                    updated_at=row[11],
                    embedding=embedding_list
                )
                entries.append(entry)

            return entries

    def get_documents(self) -> List[DocumentMetadata]:
        """Get all ingested documents"""
        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM documents ORDER BY ingested_at DESC")
            rows = cursor.fetchall()

            documents = []
            for row in rows:
                authors_list = json.loads(row[4]) if row[4] else []
                keywords_list = json.loads(row[6]) if row[6] else []

                doc = DocumentMetadata(
                    filename=row[0],
                    filepath=row[1],
                    document_type=DocumentType(row[2]),
                    title=row[3],
                    authors=authors_list,
                    abstract=row[5],
                    keywords=keywords_list,
                    publication_date=row[7],
                    doi=row[8],
                    pages=row[9],
                    ingested_at=row[10]
                )
                documents.append(doc)

            return documents

    def _on_processing_complete(self, job_id: str, metadata: DocumentMetadata, entries: List[KnowledgeEntry]):
        """Handle completed document processing"""
        try:
            with sqlite3.connect(self.db_path) as conn:
                cursor = conn.cursor()

                # Insert document metadata
                cursor.execute('''
                    INSERT OR REPLACE INTO documents
                    (filename, filepath, document_type, title, authors, abstract, keywords,
                     publication_date, doi, pages, ingested_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (
                    metadata.filename,
                    metadata.filepath,
                    metadata.document_type.value,
                    metadata.title,
                    json.dumps(metadata.authors),
                    metadata.abstract,
                    json.dumps(metadata.keywords),
                    metadata.publication_date,
                    metadata.doi,
                    metadata.pages,
                    metadata.ingested_at
                ))

                # Insert knowledge entries
                for entry in entries:
                    cursor.execute('''
                        INSERT OR REPLACE INTO knowledge_entries
                        (id, title, content, knowledge_type, document_type, source_file,
                         page_number, section, tags, metadata, created_at, updated_at, embedding)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (
                        entry.id,
                        entry.title,
                        entry.content,
                        entry.knowledge_type.value,
                        entry.document_type.value,
                        entry.source_file,
                        entry.page_number,
                        entry.section,
                        json.dumps(entry.tags),
                        json.dumps(entry.metadata),
                        entry.created_at,
                        entry.updated_at,
                        json.dumps(entry.embedding) if entry.embedding else None
                    ))

                conn.commit()

            # Clean up processor
            if job_id in self.processors:
                del self.processors[job_id]

        except Exception as e:
            print(f"Error saving processed document: {e}")

    def _on_processing_error(self, job_id: str, error: str):
        """Handle processing errors"""
        print(f"Document processing error for job {job_id}: {error}")

        # Clean up processor
        if job_id in self.processors:
            del self.processors[job_id]

class KnowledgeBaseWidget(QWidget):
    """GUI widget for knowledge base management"""

    def __init__(self, knowledge_base: KnowledgeBase, parent=None):
        super().__init__(parent)
        self.knowledge_base = knowledge_base
        self.init_ui()

    def init_ui(self):
        """Initialize the user interface"""
        layout = QVBoxLayout()

        # Title
        title = QLabel("Knowledge Base - Research Paper Ingestion")
        title.setStyleSheet("font-weight: bold; font-size: 14px;")
        layout.addWidget(title)

        # Control buttons
        button_layout = QHBoxLayout()

        self.ingest_button = QPushButton("Ingest Document")
        self.ingest_button.clicked.connect(self._ingest_document)
        button_layout.addWidget(self.ingest_button)

        self.refresh_button = QPushButton("Refresh")
        self.refresh_button.clicked.connect(self._refresh_documents)
        button_layout.addWidget(self.refresh_button)

        layout.addLayout(button_layout)

        # Search area
        search_group = QGroupBox("Search Knowledge Base")
        search_layout = QVBoxLayout()

        # Search input
        search_input_layout = QHBoxLayout()
        search_input_layout.addWidget(QLabel("Query:"))
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("Enter search terms...")
        self.search_input.returnPressed.connect(self._perform_search)
        search_input_layout.addWidget(self.search_input)

        self.search_button = QPushButton("Search")
        self.search_button.clicked.connect(self._perform_search)
        search_input_layout.addWidget(self.search_button)

        search_layout.addLayout(search_input_layout)

        # Knowledge type filter
        type_layout = QHBoxLayout()
        type_layout.addWidget(QLabel("Type:"))
        self.type_combo = QComboBox()
        self.type_combo.addItem("All Types")
        for kt in KnowledgeType:
            self.type_combo.addItem(kt.value.title())
        type_layout.addWidget(self.type_combo)
        search_layout.addLayout(type_layout)

        search_group.setLayout(search_layout)
        layout.addWidget(search_group)

        # Splitter for documents and results
        splitter = QSplitter()

        # Documents list
        doc_group = QGroupBox("Ingested Documents")
        doc_layout = QVBoxLayout()
        self.doc_list = QListWidget()
        self.doc_list.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        self.doc_list.itemDoubleClicked.connect(self._open_document)
        doc_layout.addWidget(self.doc_list)
        doc_group.setLayout(doc_layout)
        splitter.addWidget(doc_group)

        # Search results
        results_group = QGroupBox("Search Results")
        results_layout = QVBoxLayout()
        self.results_list = QListWidget()
        self.results_list.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        self.results_list.itemDoubleClicked.connect(self._show_entry_details)
        results_layout.addWidget(self.results_list)
        results_group.setLayout(results_layout)
        splitter.addWidget(results_group)

        layout.addWidget(splitter)

        # Progress bar for ingestion
        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        layout.addWidget(self.progress_bar)

        # Status label
        self.status_label = QLabel("Ready")
        layout.addWidget(self.status_label)

        self.setLayout(layout)

        # Set size policies for resizability
        self.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        self.setMinimumSize(400, 300)  # Minimum size to ensure usability

        # Initial refresh
        self._refresh_documents()

    def _ingest_document(self):
        """Ingest a new document"""
        from PyQt6.QtWidgets import QFileDialog

        file_dialog = QFileDialog()
        file_dialog.setFileMode(QFileDialog.FileMode.ExistingFile)
        file_dialog.setNameFilter("Documents (*.pdf *.docx *.txt)")

        if file_dialog.exec():
            filepath = file_dialog.selectedFiles()[0]
            if filepath:
                self._start_ingestion(filepath)

    def _start_ingestion(self, filepath: str):
        """Start document ingestion process"""
        self.progress_bar.setVisible(True)
        self.progress_bar.setRange(0, 0)  # Indeterminate progress
        self.status_label.setText(f"Ingesting: {os.path.basename(filepath)}")

        # Create processor and connect signals
        processor = DocumentProcessor(filepath)
        processor.processing_complete.connect(self._on_ingestion_complete)
        processor.processing_progress.connect(self._on_ingestion_progress)
        processor.processing_error.connect(self._on_ingestion_error)

        processor.start()

    def _on_ingestion_complete(self, metadata: DocumentMetadata, entries: List[KnowledgeEntry]):
        """Handle completed ingestion"""
        self.progress_bar.setVisible(False)
        self.status_label.setText(f"Successfully ingested: {metadata.filename} ({len(entries)} entries)")
        self._refresh_documents()

    def _on_ingestion_progress(self, status: str, progress: float):
        """Handle ingestion progress updates"""
        self.status_label.setText(status)
        if progress >= 0:
            self.progress_bar.setRange(0, 100)
            self.progress_bar.setValue(int(progress * 100))

    def _on_ingestion_error(self, error: str):
        """Handle ingestion errors"""
        self.progress_bar.setVisible(False)
        self.status_label.setText(f"Ingestion failed: {error}")

    def _refresh_documents(self):
        """Refresh the documents list"""
        self.doc_list.clear()
        documents = self.knowledge_base.get_documents()

        for doc in documents:
            item_text = f"{doc.filename}"
            if doc.title:
                item_text += f" - {doc.title}"
            if doc.authors:
                item_text += f" ({', '.join(doc.authors)})"

            item = QListWidgetItem(item_text)
            item.setData(1, doc)  # Store document metadata
            self.doc_list.addItem(item)

    def _open_document(self, item: QListWidgetItem):
        """Open the selected document"""
        doc = item.data(1)
        if doc and os.path.exists(doc.filepath):
            QDesktopServices.openUrl(QUrl.fromLocalFile(doc.filepath))

    def _perform_search(self):
        """Perform knowledge base search"""
        query = self.search_input.text().strip()
        if not query:
            return

        # Get knowledge type filter
        type_index = self.type_combo.currentIndex()
        knowledge_types = None
        if type_index > 0:  # Not "All Types"
            kt = list(KnowledgeType)[type_index - 1]
            knowledge_types = [kt]

        # Perform search
        results = self.knowledge_base.search_knowledge(query, knowledge_types)

        # Update results list
        self.results_list.clear()
        for entry in results:
            item_text = f"[{entry.knowledge_type.value}] {entry.title}"
            if entry.section:
                item_text += f" ({entry.section})"
            item_text += f" - {entry.source_file}"

            item = QListWidgetItem(item_text)
            item.setData(1, entry)  # Store entry data
            self.results_list.addItem(item)

        self.status_label.setText(f"Found {len(results)} results for '{query}'")

    def _show_entry_details(self, item: QListWidgetItem):
        """Show detailed information for a knowledge entry"""
        entry = item.data(1)
        if entry:
            from PyQt6.QtWidgets import QDialog, QVBoxLayout, QTextEdit, QLabel

            dialog = QDialog(self)
            dialog.setWindowTitle(f"Knowledge Entry: {entry.title}")
            dialog.resize(600, 400)

            layout = QVBoxLayout()

            # Metadata
            meta_text = f"Type: {entry.knowledge_type.value}\n"
            meta_text += f"Source: {entry.source_file}\n"
            if entry.section:
                meta_text += f"Section: {entry.section}\n"
            if entry.tags:
                meta_text += f"Tags: {', '.join(entry.tags)}\n"
            meta_text += f"Created: {datetime.fromtimestamp(entry.created_at).strftime('%Y-%m-%d %H:%M')}"

            meta_label = QLabel(meta_text)
            meta_label.setStyleSheet("font-weight: bold;")
            layout.addWidget(meta_label)

            # Content
            content_edit = QTextEdit()
            content_edit.setPlainText(entry.content)
            content_edit.setReadOnly(True)
            layout.addWidget(content_edit)

            dialog.setLayout(layout)
            dialog.exec()

# Export main classes
__all__ = ['KnowledgeBase', 'KnowledgeBaseWidget', 'KnowledgeEntry', 'DocumentMetadata', 'DocumentType', 'KnowledgeType']